"""Approval form generation and reimbursement cheat sheet for university travel."""

import os
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


def calculate_reimbursement_summary(
    trip_facts: dict[str, Any],
    transport_cost: float = 0.0,
    lodging_cost: float = 0.0,
) -> dict[str, Any]:
    """Calculate allowances and budget projection according to university policy.
    
    Rule:
    - Meal allowance: 100 RMB/day (ordinary) or 120 RMB/day (Tibet/Qinghai/Xinjiang).
    - Local transport allowance: 80 RMB/day flat package.
    - Calculated based on natural calendar days.
    """
    dep_str = trip_facts.get("departure_date")
    ret_str = trip_facts.get("return_date") or trip_facts.get("departure_date")
    destination = trip_facts.get("destination", "")

    total_days = 1
    if dep_str and ret_str:
        try:
            d1 = date.fromisoformat(str(dep_str))
            d2 = date.fromisoformat(str(ret_str))
            total_days = max(1, (d2 - d1).days + 1)
        except (ValueError, TypeError):
            total_days = 1

    # Special allowance regions
    special_regions = {"西藏", "青海", "新疆", "拉萨", "西宁", "乌鲁木齐"}
    daily_meal = 120.0 if any(r in destination for r in special_regions) else 100.0
    daily_local_transport = 80.0

    total_meal = daily_meal * total_days
    total_local_transport = daily_local_transport * total_days
    total_budget = transport_cost + lodging_cost + total_meal + total_local_transport

    return {
        "days": total_days,
        "daily_meal": daily_meal,
        "total_meal": total_meal,
        "daily_local_transport": daily_local_transport,
        "total_local_transport": total_local_transport,
        "transport_cost": transport_cost,
        "lodging_cost": lodging_cost,
        "total_budget": total_budget,
    }


def render_reimbursement_markdown(
    trip_facts: dict[str, Any],
    transport_cost: float = 0.0,
    lodging_cost: float = 0.0,
) -> str:
    """Render a structured markdown cheat sheet for filing expense report in university portal."""
    summary = calculate_reimbursement_summary(trip_facts, transport_cost, lodging_cost)
    applicant = trip_facts.get("user_name") or "出差人"
    origin = trip_facts.get("origin") or "常驻地"
    dest = trip_facts.get("destination") or "出差地"
    dates = f"{trip_facts.get('departure_date', '')} 至 {trip_facts.get('return_date') or trip_facts.get('departure_date', '')}"
    purpose = trip_facts.get("purpose") or "因公出差 / 学术交流"

    lines = [
        "### 📋 合肥工业大学差旅报销预审对账单（系统填报参考）",
        f"> **出差人员**：{applicant} | **出差事由**：{purpose}",
        f"> **行程路线**：{origin} ➔ {dest} | **出差日期**：{dates}（共 {summary['days']} 个自然天）",
        "",
        "**预估费用与包干补助明细**（登录学校财务系统时按此项填报）：",
        f"1. **城市间交通费**：¥{summary['transport_cost']:.2f}",
        f"2. **住宿费**：¥{summary['lodging_cost']:.2f}",
        f"3. **伙食补助费（定额包干）**：¥{summary['daily_meal']:.0f}/天 × {summary['days']}天 = **¥{summary['total_meal']:.2f}**",
        f"4. **市内交通补助费（定额包干）**：¥{summary['daily_local_transport']:.0f}/天 × {summary['days']}天 = **¥{summary['total_local_transport']:.2f}**",
        "",
        f"💰 **报销总预算预估**：**¥{summary['total_budget']:.2f}**",
        "",
        "**财务报销合规提醒**：",
        "- 凭据报销：请保留机票行程单/登机牌、高铁车票、住宿发票（发票须注明单价与间夜数）；",
        "- 时限与结算：出差结束后两个月内报销，须使用公务卡或银行对公转账结算；",
        "- 事前审批：报销时需附上已由单位负责人签字的《合肥工业大学出差审批单》。",
    ]
    return "\n".join(lines)


def build_approval_sheet_doc(
    trip_facts: dict[str, Any],
    *,
    is_rental_or_driving: bool = False,
    budget_estimate: float = 0.0,
) -> Document:
    """Build and return a python-docx Document object in memory."""
    doc = Document()

    applicant = trip_facts.get("user_name") or "出差人员"
    dept = trip_facts.get("department") or "相关学院/科研单位"
    dest = trip_facts.get("destination") or "目的地"
    origin = trip_facts.get("origin") or "合肥"
    route = f"{origin} - {dest} - {origin}"
    purpose = trip_facts.get("purpose") or "因公开展教学科研学术活动"
    dep_date = str(trip_facts.get("departure_date") or date.today())
    ret_date = str(trip_facts.get("return_date") or dep_date)
    duration = f"{dep_date} 至 {ret_date}"

    # Document Title
    title_text = "合肥工业大学租车（自驾车）出差审批表" if is_rental_or_driving else "合肥工业大学出差审批单"
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(title_text)
    run.font.name = "SimSun"
    run.font.size = Pt(16)
    run.font.bold = True

    # Department Subheader
    unit_p = doc.add_paragraph()
    unit_run = unit_p.add_run(f"单位：{dept}")
    unit_run.font.name = "SimSun"
    unit_run.font.size = Pt(11)

    if not is_rental_or_driving:
        # Standard Travel Approval Table (5 rows, 2 columns)
        table = doc.add_table(rows=5, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        fields = [
            ("出差人", applicant),
            ("出差时间", duration),
            ("出差路线", route),
            ("出差事由", purpose),
            ("审批意见", "\n\n签字：                      年   月   日\n"),
        ]
        for row_idx, (label, val) in enumerate(fields):
            row = table.rows[row_idx]
            c0, c1 = row.cells[0], row.cells[1]
            c0.width = Inches(1.5)
            c1.width = Inches(5.0)
            c0.text = label
            c1.text = val
            for cell in (c0, c1):
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(6)
                    p.paragraph_format.space_after = Pt(6)
                    for r in p.runs:
                        r.font.name = "SimSun"
                        r.font.size = Pt(11)
    else:
        # Rental / Driving Approval Table (6 rows)
        table = doc.add_table(rows=6, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        r0 = table.rows[0]
        r0.cells[0].text, r0.cells[1].text = "申请人", applicant
        r0.cells[2].text, r0.cells[3].text = "工作部门", dept

        r1 = table.rows[1]
        r1.cells[0].text, r1.cells[1].text = "财务号", trip_facts.get("finance_code", "")
        r1.cells[2].text, r1.cells[3].text = "预算金额", f"¥{budget_estimate:.2f}" if budget_estimate else ""

        r2 = table.rows[2]
        r2.cells[0].text = "出差事由"
        r2.cells[1].merge(r2.cells[3]).text = purpose

        r3 = table.rows[3]
        r3.cells[0].text = "地点及起止时间"
        r3.cells[1].merge(r3.cells[3]).text = f"地点：{dest}\n时间：{duration}"

        r4 = table.rows[4]
        r4.cells[0].text = "交通说明"
        r4.cells[1].merge(r4.cells[3]).text = trip_facts.get("vehicle_detail", "租车/自驾出行合规说明")

        r5 = table.rows[5]
        r5.cells[0].text = "单位（项目）负责人审批"
        r5.cells[1].merge(r5.cells[3]).text = "\n\n签字：                      年   月   日\n"

        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(4)
                    for r in p.runs:
                        r.font.name = "SimSun"
                        r.font.size = Pt(10.5)

        note_p = doc.add_paragraph()
        note_run = note_p.add_run("注：因自驾车或者租车所引起的安全等问题，由出差人员承担。")
        note_run.font.name = "SimSun"
        note_run.font.size = Pt(9.5)
        note_run.font.italic = True

    return doc


def export_approval_sheet_docx(
    trip_facts: dict[str, Any],
    output_path: str | Path | None = None,
    *,
    is_rental_or_driving: bool = False,
    budget_estimate: float = 0.0,
) -> Path | BytesIO:
    """Generate a clean Word document of Hefei University of Technology travel approval form.
    
    If output_path is provided, writes to file. Otherwise returns BytesIO.
    """
    doc = build_approval_sheet_doc(
        trip_facts,
        is_rental_or_driving=is_rental_or_driving,
        budget_estimate=budget_estimate,
    )
    if output_path is not None:
        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out))
        return out

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
