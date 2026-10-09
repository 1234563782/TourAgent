"""Offline policy verification: native Word, rule boundaries, embeddings and scoped SQL."""

import json
import os
import tempfile
import unittest
from dataclasses import replace
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import MagicMock, Mock, patch

import httpx
from docx import Document

from business_travel.contracts import PolicyRules
from business_travel.policy import PolicyStore
from business_travel.policy_documents import PolicyBundle, PolicyChunk, parse_docx
from business_travel.policy_pg import (
    ChineseTokenizer, EmbeddingClient, PolicyError, PostgresPolicyStore, reciprocal_rank_fusion, validate_vector,
)
from business_travel.policy_rules import ExpenseFacts, ReviewedRules, check_expense
from business_travel.settings import Settings
from scripts.policy_admin import main

ROOT = Path(__file__).resolve().parents[1]


def settings(**changes):
    base = Settings(api_key="", model="", base_url="", node_path="", flyai_cli_path="",
                    storage_dir=Path("."), legacy_memory_dir=Path("."),
                    policy_backend="pgvector", policy_company_id="test_co", policy_key="travel",
                    policy_database_url="postgresql://synthetic-secret",
                    embedding_base_url="https://embedding.example/v1", embedding_api_key="test-secret",
                    embedding_model="test-embedding", embedding_dimensions=3)
    return replace(base, **changes)


def bundle():
    return PolicyBundle(company_id="test_co", policy_key="travel", version="test",
                        effective_from="2026-09-07", source="synthetic.docx", source_sha256="a" * 64,
                        chunks=[PolicyChunk(id="one", clause="4.2", section="住宿", text="普通员工住宿420元",
                                            related_ids=["two"]),
                                PolicyChunk(id="two", clause="4.3", section="住宿", text="合住需核对")])


def database_doc(**changes):
    return {
        "id": 1, "company_id": "test_co", "policy_key": "travel", "version": "test",
        "effective_from": date(2026, 9, 7), "effective_to": None, "source": "synthetic.docx",
        "source_sha256": "a" * 64, "bundle_sha256": bundle().digest(), "status": "draft",
        "embedding_space": "test-space", "tokenizer": "test-tokenizer", **changes,
    }


def chunk_row(key, related=()):
    return {"chunk_id": key, "clause": "4.2", "section": "住宿", "content": "模拟政策",
            "table_data": None, "related_ids": list(related)}


def store_with_db(rows):
    connection = MagicMock()
    connection.__enter__.return_value = connection
    answers = iter(rows)

    def execute(sql, params=None):
        cursor = Mock()
        if sql.startswith("SELECT set_config") or sql.startswith("SELECT pg_advisory"):
            return cursor
        data = next(answers)
        cursor.fetchone.return_value = data
        cursor.fetchall.return_value = data
        return cursor

    connection.execute.side_effect = execute
    connect = Mock(return_value=connection)
    embedder = Mock(space="test-space")
    embedder.embed.side_effect = lambda texts: [[1.0, 0.0, 0.0] for _ in texts]
    tokenizer = Mock(signature="test-tokenizer")
    tokenizer.tokenize.return_value = "住宿 420"
    store = PostgresPolicyStore(settings(), connect=connect, embedder=embedder, tokenizer=tokenizer)
    return store, connection, connect


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "policy.docx"

    def parse(self, **kwargs):
        return parse_docx(self.path, company_id="test_co", policy_key="travel",
                          version="test", effective_from=date(2026, 9, 7), **kwargs)

    def make_doc(self):
        doc = Document()
        doc.add_heading("第四章 住宿", 1)
        doc.add_paragraph("4.2 每间每晚含税标准")
        table = doc.add_table(rows=2, cols=2)
        for cell, text in zip(table.rows[0].cells, ["职级", "A类 元/间/晚"]):
            cell.text = text
        for cell, text in zip(table.rows[1].cells, ["普通员工", "420"]):
            cell.text = text
        doc.add_paragraph("4.3 合住按4.2计算")
        doc.save(self.path)
        return doc

    def test_table_order_headers_source_and_links(self):
        self.make_doc()
        before = self.path.read_bytes()
        result = self.parse()
        row = next(c for c in result.chunks if c.kind == "table_row")
        self.assertIn("A类 元/间/晚：420", row.text)
        self.assertEqual(row.table["cells"], ["普通员工", "420"])
        self.assertEqual(row.clause, "4.2")
        self.assertTrue(any(c.clause == "4.3" and c.id in row.related_ids for c in result.chunks))
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(result.digest(), self.parse().digest())
        self.assertIsNone(result.effective_to)

    def test_merged_grid_warns_and_preserves(self):
        doc = self.make_doc()
        doc.tables[0].rows[0].cells[0].merge(doc.tables[0].rows[0].cells[1])
        doc.save(self.path)
        result = self.parse()
        self.assertTrue(any("合并" in w for w in result.warnings))
        self.assertEqual(len(next(c for c in result.chunks if c.table).table["headers"]), 2)

    def test_nested_table_rejected_not_silently_lost(self):
        doc = self.make_doc()
        doc.tables[0].cell(1, 0).add_table(rows=1, cols=1)
        doc.save(self.path)
        with self.assertRaises(ValueError):
            self.parse()

    def test_empty_document_rejected(self):
        Document().save(self.path)
        with self.assertRaises(ValueError):
            self.parse()

    def test_long_clause_parts_link_together(self):
        doc = Document()
        doc.add_paragraph("4.2 " + "住宿标准。" * 120)
        doc.save(self.path)
        result = self.parse(chunk_size=200)
        self.assertGreater(len(result.chunks), 1)
        self.assertTrue(all(c.clause == "4.2" for c in result.chunks))
        self.assertIn(result.chunks[1].id, result.chunks[0].related_ids)

    def test_duplicate_ids_invalid_references_and_dates_rejected(self):
        value = bundle().model_dump(mode="json")
        for update in (
            {"chunks": [value["chunks"][0], value["chunks"][0]]},
            {"chunks": [value["chunks"][0]]},
            {"effective_to": "2020-01-01"},
            {"company_id": "../other"},
        ):
            with self.assertRaises(ValueError):
                PolicyBundle.model_validate({**value, **update})

    def test_prepare_does_not_load_env_or_call_network_and_refuses_overwrite(self):
        self.make_doc()
        output = Path(self.temp.name) / "bundle.json"
        args = ["prepare", "--source", str(self.path), "--output", str(output),
                "--company", "test_co", "--version", "test", "--effective-from", "2026-09-07"]
        with patch.object(Settings, "from_env", side_effect=AssertionError("must stay offline")), patch("builtins.print"):
            self.assertEqual(main(args), 0)
            before = output.read_bytes()
            self.assertEqual(main(args), 2)
            self.assertEqual(output.read_bytes(), before)


class RuleTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.rules = ReviewedRules.model_validate_json((ROOT / "policies/demo-v1.1.rules.json").read_text("utf-8"))

    def facts(self, **changes):
        return ExpenseFacts(kind="hotel", expense_date="2026-09-09", destination="上海",
                            **{"employee_grade": "普通员工", "mainland_business_trip": True,
                               "long_assignment_days": 2, "special_approval": False, "amount": "400",
                               "currency": "CNY", "basis": "room_night", "taxes_confirmed": True,
                               "date_verified": True, "occupants": 1, **changes})

    def check(self, **changes):
        return check_expense(self.facts(**changes), self.rules)

    def test_currency_contract_rejects_names_and_preserves_explicit_codes(self):
        from pydantic import ValidationError
        from business_travel.slots import ExpensePatch
        for value in ("人民币", "美元", "cny", " CNY ", "$", ""):
            with self.subTest(currency=value):
                with self.assertRaises(ValidationError):
                    self.facts(currency=value)
                with self.assertRaises(ValidationError):
                    ExpensePatch(currency=value)
        for value in ("CNY", "USD", "EUR", "JPY", None):
            with self.subTest(currency=value):
                self.assertEqual(self.facts(currency=value).currency, value)
                self.assertEqual(ExpensePatch(currency=value).currency, value)
        self.assertEqual(self.check(currency="USD")["status"], "insufficient_information")
        self.assertEqual(self.check(currency=None)["status"], "insufficient_information")

    def test_hotel_below_equal_above(self):
        for amount, expected in (("400", "within_checked_rules"), ("420", "within_checked_rules"),
                                 ("420.01", "over_standard")):
            with self.subTest(amount=amount):
                result = self.check(amount=amount)
                self.assertEqual(result["status"], expected)
                self.assertEqual(result["cap"], "420")
                self.assertFalse(result["approved"])

    def test_shared_room_uses_higher_cap_times_multiplier(self):
        result = self.check(occupants=2, roommate_grade="经理主管", voluntary_twin_share=True, amount="700")
        self.assertEqual(Decimal(result["cap"]), Decimal("750"))
        self.assertEqual(result["status"], "within_checked_rules")

    def test_special_approval_routes_to_manual_without_unrelated_scope_facts(self):
        result = self.check(special_approval=True, mainland_business_trip=None, long_assignment_days=None)
        self.assertEqual(result["status"], "manual_review")
        self.assertEqual(result["reason"], "verify_special_approval_scope")
        self.assertFalse(result["approved"])

    def test_unknowns_never_pass(self):
        for values in ({"employee_grade": None}, {"occupants": None}, {"taxes_confirmed": False},
                       {"date_verified": False}, {"basis": "unknown"}, {"amount": None},
                       {"currency": "USD"}, {"mainland_business_trip": None},
                       {"long_assignment_days": None}, {"special_approval": None},
                       {"occupants": 2, "roommate_grade": None}):
            with self.subTest(values=values):
                self.assertEqual(self.check(**values)["status"], "insufficient_information")

    def test_city_aliases_and_unknown_city(self):
        facts = self.facts().model_copy(update={"destination": "上海市"})
        self.assertEqual(check_expense(facts, self.rules)["cap"], "420")
        facts = facts.model_copy(update={"destination": "浦东新区"})
        self.assertEqual(check_expense(facts, self.rules)["reason"], "city_mapping_required")

    def test_long_assignment_exact_boundary(self):
        self.assertEqual(self.check(long_assignment_days=29)["status"], "within_checked_rules")
        for days in (30, 31):
            self.assertEqual(self.check(long_assignment_days=days)["status"], "manual_review")

    def test_special_approval_is_not_automatic_pass(self):
        self.assertEqual(self.check(special_approval=True)["status"], "manual_review")

    def test_non_mainland_or_special_room_requires_review(self):
        self.assertEqual(self.check(mainland_business_trip=False)["status"], "manual_review")
        self.assertEqual(self.check(occupants=3)["status"], "manual_review")

    def meal(self, **changes):
        base = self.facts().model_dump()
        base.update(kind="meal", overnight_trip=True, allowance_region="普通地区",
                    covered_meals=["breakfast", "lunch"])
        return check_expense(ExpenseFacts.model_validate({**base, **changes}), self.rules)

    def test_meal_examples_and_deduplication(self):
        self.assertEqual(Decimal(self.meal()["amount"]), Decimal("40"))
        self.assertEqual(Decimal(self.meal(covered_meals=["breakfast", "lunch", "lunch"])["amount"]), Decimal("40"))
        self.assertEqual(Decimal(self.meal(overnight_trip=False, covered_meals=["lunch"])["amount"]), Decimal("30"))
        self.assertEqual(Decimal(self.meal(allowance_region="西藏", covered_meals=[])["amount"]), Decimal("120"))
        self.assertEqual(Decimal(self.meal(covered_meals=["breakfast", "lunch", "dinner"])["amount"]), Decimal("0"))

    def test_meal_missing_conditions(self):
        for values in ({"overnight_trip": None}, {"covered_meals": None},
                       {"allowance_region": None}, {"allowance_region": "不明地区"}):
            self.assertEqual(self.meal(**values)["status"], "insufficient_information")

    def local(self, **changes):
        base = self.facts().model_dump()
        base.update(kind="local_transport", local_business_travel=True, local_transport_coverage="partial",
                    self_paid_local_travel=True, local_car_cost_reimbursed=False)
        return check_expense(ExpenseFacts.model_validate({**base, **changes}), self.rules)

    def test_local_transport_conditions(self):
        self.assertEqual(self.local()["amount"], "80")
        for values in ({"local_business_travel": False}, {"local_transport_coverage": "all"},
                       {"local_car_cost_reimbursed": True}, {"self_paid_local_travel": False}):
            self.assertEqual(self.local(**values)["amount"], "0")
        self.assertEqual(self.local(local_business_travel=None)["status"], "insufficient_information")

    def transport(self, **changes):
        base = self.facts().model_dump()
        base.update(kind="transport", transport_mode="flight", seat_class="经济舱",
                    flight_approval=True, suitable_direct_train_hours="7")
        return check_expense(ExpenseFacts.model_validate({**base, **changes}), self.rules)

    def test_transport_priority_and_approval(self):
        self.assertEqual(self.transport()["status"], "within_checked_rules")
        self.assertEqual(self.transport(flight_approval=None)["status"], "manual_review")
        self.assertEqual(self.transport(seat_class="公务舱")["status"], "over_standard")
        self.assertEqual(self.transport(suitable_direct_train_hours="6")["status"], "manual_review")
        self.assertEqual(self.transport(suitable_direct_train_hours=None)["status"], "insufficient_information")
        self.assertEqual(self.transport(transport_mode="high_speed_train", seat_class="二等座")["status"],
                         "within_checked_rules")

    def test_invalid_rule_values_fail(self):
        for values in ({"meal_deductions": {"breakfast": "1", "lunch": "1", "dinner": "1"}},
                       {"source_clauses": {}}, {"hotel_caps": {"普通员工": {"A": "-1"}}},
                       {"meal_region_bases": {"普通地区": "NaN"}}):
            with self.assertRaises(ValueError):
                ReviewedRules.model_validate({**self.rules.model_dump(), **values})

    def test_unbounded_legacy_version_supported(self):
        rules = PolicyRules(version="test", effective_from="2026-09-07")
        self.assertIsNone(rules.effective_to)


class EmbeddingTests(unittest.TestCase):
    def test_response_reordered_by_index(self):
        def handler(request):
            body = json.loads(request.content)
            self.assertEqual(body["model"], "test-embedding")
            self.assertNotIn("dimensions", body)
            return httpx.Response(200, json={"data": [
                {"index": 1, "embedding": [0, 1, 0]}, {"index": 0, "embedding": [1, 0, 0]}]})
        client = EmbeddingClient(settings(), transport=httpx.MockTransport(handler))
        self.assertEqual(client.embed(["one", "two"]), [[1., 0., 0.], [0., 1., 0.]])

    def test_vector_validation(self):
        for vector in ([0, 0, 0], [1, 2], [True, 0, 1], [float("nan"), 1, 0], "invalid"):
            with self.assertRaises(PolicyError):
                validate_vector(vector, 3)

    def test_rejected_responses_never_expose_secret(self):
        for status in (301, 401, 429, 500):
            transport = httpx.MockTransport(lambda request: httpx.Response(status, text="test-secret"))
            with self.assertRaises(PolicyError) as error:
                EmbeddingClient(settings(), transport=transport).embed(["one"])
            self.assertNotIn("test-secret", str(error.exception))

    def test_duplicate_indices_rejected(self):
        transport = httpx.MockTransport(lambda request: httpx.Response(200, json={"data": [
            {"index": 0, "embedding": [1, 0, 0]}, {"index": 0, "embedding": [0, 1, 0]}]}))
        with self.assertRaises(PolicyError):
            EmbeddingClient(settings(), transport=transport).embed(["one", "two"])

    def test_untrusted_base_url_never_called(self):
        handler = Mock()
        for url in ("http://example.com", "https://key@example.com", "https://example.com?key=x"):
            with self.assertRaises(PolicyError):
                EmbeddingClient(settings(embedding_base_url=url), transport=httpx.MockTransport(handler)).embed(["one"])
        handler.assert_not_called()

    def test_space_changes_with_provider_model_revision_dimensions(self):
        client = EmbeddingClient(settings())
        for changes in ({"embedding_model": "other"}, {"embedding_revision": "2"},
                        {"embedding_dimensions": 4}, {"embedding_base_url": "https://other.example/v1"}):
            self.assertNotEqual(client.space, EmbeddingClient(settings(**changes)).space)

    def test_chinese_terms_and_rrf(self):
        tokenizer = ChineseTokenizer()
        tokens = tokenizer.tokenize("上海普通员工住宿420元，合住？")
        self.assertIn("上海", tokens)
        self.assertIn("420", tokens)
        self.assertNotIn("？", tokens)
        self.assertEqual(reciprocal_rank_fusion(["a", "b"], ["b", "c"], limit=1), ["b"])


class RepositoryTests(unittest.TestCase):
    def test_stage_defaults_to_single_input_requests(self):
        store, _, _ = store_with_db([None, {"id": 9}, None, None])
        store.stage(bundle())
        self.assertEqual([len(c.args[0]) for c in store.embedder.embed.call_args_list], [1, 1])

    def test_stage_respects_configured_batch_size(self):
        store, _, _ = store_with_db([None, {"id": 9}, None, None])
        store.settings = settings(embedding_batch_size=2)
        store.stage(bundle())
        self.assertEqual([len(c.args[0]) for c in store.embedder.embed.call_args_list], [2])

    def test_batch_size_environment_validation(self):
        for value in ("0", "65", "-1", "invalid"):
            with self.subTest(value=value), patch.dict(
                    os.environ, {"TRAVEL_LOAD_DOTENV": "0", "TRAVEL_EMBEDDING_BATCH_SIZE": value}, clear=True):
                with self.assertRaises(ValueError):
                    Settings.from_env()
        with patch.dict(os.environ, {"TRAVEL_LOAD_DOTENV": "0",
                                    "TRAVEL_EMBEDDING_BATCH_SIZE": "2"}, clear=True):
            self.assertEqual(Settings.from_env().embedding_batch_size, 2)

    def test_missing_database_config_never_connects(self):
        connect = Mock()
        store = PostgresPolicyStore(settings(policy_database_url=""), connect=connect)
        self.assertEqual(store.search("住宿")["status"], "unavailable")
        connect.assert_not_called()

    def test_scope_mismatch_never_embeds(self):
        store, _, _ = store_with_db([])
        with self.assertRaises(PolicyError):
            store.stage(bundle().model_copy(update={"company_id": "other_co"}))
        store.embedder.embed.assert_not_called()

    def test_idempotent_stage_does_not_embed(self):
        store, _, _ = store_with_db([database_doc()])
        self.assertEqual(store.stage(bundle())["status"], "unchanged")
        store.embedder.embed.assert_not_called()

    def test_conflicting_version_refused(self):
        store, _, _ = store_with_db([database_doc(bundle_sha256="b" * 64)])
        with self.assertRaises(PolicyError):
            store.stage(bundle())
        store.embedder.embed.assert_not_called()

    def test_stage_writes_whole_document_as_draft_transaction(self):
        store, conn, _ = store_with_db([None, {"id": 9}, None, None])
        result = store.stage(bundle())
        self.assertEqual(result["status"], "draft")
        inserts = [call.args for call in conn.execute.call_args_list if call.args[0].startswith("INSERT")]
        self.assertEqual(len(inserts), 3)
        self.assertNotIn("published", inserts[0][0])
        self.assertEqual(inserts[1][1][0], 9)

    def test_stage_failure_exits_transaction_with_error(self):
        store, conn, _ = store_with_db([None, {"id": 9}])
        with self.assertRaises(PolicyError):
            store.stage(bundle())
        self.assertIsNotNone(conn.__exit__.call_args.args[0])

    def test_no_published_version_does_not_embed(self):
        store, _, _ = store_with_db([[]])
        self.assertEqual(store.search("住宿")["status"], "empty")
        store.embedder.embed.assert_not_called()

    def test_search_current_publication_scope_and_context(self):
        store, conn, _ = store_with_db([
            [database_doc(status="published")], [{"chunk_id": "one"}], [{"chunk_id": "one"}],
            [chunk_row("one", ["two"])], [chunk_row("two")],
        ])
        result = store.search("住宿';DROP TABLE x;")
        self.assertEqual([row["chunk"] for row in result["sources"]], ["one", "two"])
        active = conn.execute.call_args_list[1]
        self.assertEqual(active.args[1], ("test_co", "travel"))
        self.assertIn("status='published'", active.args[0])
        for call in conn.execute.call_args_list:
            self.assertNotIn("DROP TABLE", call.args[0])
            if "policy_chunks" in call.args[0]:
                self.assertIn("document_id=%s", call.args[0])
        self.assertFalse(result["context_incomplete"])

    def test_search_adds_reviewed_scope_without_related_links(self):
        boundary = {**chunk_row("boundary"), "clause": "scope.example"}
        store, conn, _ = store_with_db([
            [database_doc(status="published", scope_clauses=["scope.example"])],
            [{"chunk_id": "one"}], [{"chunk_id": "one"}],
            [chunk_row("one")], [boundary],
        ])
        result = store.search("审批边界")
        self.assertEqual([row["chunk"] for row in result["sources"]], ["one", "boundary"])
        self.assertFalse(result["context_incomplete"])
        scope_query = next(c for c in conn.execute.call_args_list
                           if "clause=ANY" in c.args[0])
        self.assertIn("document_id=%s", scope_query.args[0])
        self.assertIn("LIMIT %s", scope_query.args[0])
        self.assertEqual(scope_query.args[1][:3], (1, ["scope.example"], ["one"]))

    def test_search_scope_overflow_preserves_primary_and_reports_incomplete(self):
        store, _, _ = store_with_db([
            [database_doc(scope_clauses=["scope.example"])],
            [{"chunk_id": "one"}], [{"chunk_id": "one"}],
            [chunk_row("one")], [chunk_row("boundary")],
        ])
        store.settings = replace(store.settings, policy_context_limit=1)
        result = store.search("审批边界")
        self.assertEqual([row["chunk"] for row in result["sources"]], ["one"])
        self.assertTrue(result["context_incomplete"])

    def test_search_scope_and_links_share_context_limit(self):
        store, _, _ = store_with_db([
            [database_doc(scope_clauses=["scope.example"])],
            [{"chunk_id": "one"}], [{"chunk_id": "one"}],
            [chunk_row("one", ["two"])], [chunk_row("boundary")], [],
        ])
        store.settings = replace(store.settings, policy_context_limit=2)
        result = store.search("审批边界")
        self.assertEqual([row["chunk"] for row in result["sources"]], ["one", "boundary"])
        self.assertTrue(result["context_incomplete"])

    def test_search_scope_does_not_duplicate_already_selected_links(self):
        store, _, _ = store_with_db([
            [database_doc(scope_clauses=["scope.example"])],
            [{"chunk_id": "one"}], [{"chunk_id": "one"}],
            [chunk_row("one", ["boundary"])], [chunk_row("boundary")],
        ])
        result = store.search("审批边界")
        self.assertEqual([row["chunk"] for row in result["sources"]], ["one", "boundary"])
        self.assertFalse(result["context_incomplete"])

    def test_changed_embedding_space_refused(self):
        store, _, _ = store_with_db([[database_doc(embedding_space="different")]])
        self.assertEqual(store.search("住宿")["reason"], "policy_embedding_or_tokenizer_changed_reimport_required")
        store.embedder.embed.assert_not_called()

    def test_overlapping_versions_fail_closed(self):
        store, _, _ = store_with_db([[database_doc(), database_doc(id=2)]])
        self.assertEqual(store.search("住宿")["reason"], "ambiguous_policy_versions")

    def test_publish_needs_matching_review_hash(self):
        store, _, _ = store_with_db([database_doc()])
        with self.assertRaises(PolicyError):
            store.publish(1, "wrong")

    def test_publish_overlap_requires_explicit_supersede(self):
        store, conn, _ = store_with_db([database_doc(), [{"id": 2, "effective_from": date(2020, 1, 1)}]])
        with self.assertRaises(PolicyError):
            store.publish(1, bundle().digest())
        self.assertFalse(any(c.args[0].lstrip().startswith(("UPDATE ", "INSERT ", "DELETE "))
                             for c in conn.execute.call_args_list))

    def test_supersede_closes_old_version_only(self):
        store, conn, _ = store_with_db([
            database_doc(), [{"id": 2, "effective_from": date(2020, 1, 1)}], None, None])
        self.assertEqual(store.publish(1, bundle().digest(), supersede=True)["status"], "published")
        update = next(c for c in conn.execute.call_args_list if "SET effective_to" in c.args[0])
        self.assertEqual(update.args[1], (date(2026, 9, 6), 2))

    def test_database_errors_do_not_leak_connection_string(self):
        store, _, connect = store_with_db([])
        connect.side_effect = RuntimeError("postgresql://synthetic-secret")
        result = store.search("住宿")
        self.assertNotIn("synthetic-secret", str(result))
        self.assertEqual(result["status"], "unavailable")

    def test_rule_check_version_sources_and_redacted_audit(self):
        rules = json.loads((ROOT / "policies/demo-v1.1.rules.json").read_text("utf-8"))
        store, conn, _ = store_with_db([
            [database_doc(status="published")], {"rules": rules}, [chunk_row("one")], None])
        facts = ExpenseFacts.model_validate_json((ROOT / "policies/demo-hotel-facts.json").read_text("utf-8"))
        result = store.check(facts)
        self.assertEqual(result["status"], "within_checked_rules")
        self.assertEqual(result["policy_version"], "test")
        self.assertEqual(result["sources"][0]["source"], "synthetic.docx")
        audit = next(c for c in conn.execute.call_args_list if "INSERT INTO policy_check_logs" in c.args[0])
        data = audit.args[1][-1].obj
        self.assertEqual(set(data), {"status", "reason", "source_clauses"})
        self.assertNotIn("上海", str(data))

    def test_unknown_rules_or_missing_version_not_approved(self):
        facts = ExpenseFacts.model_validate_json((ROOT / "policies/demo-hotel-facts.json").read_text("utf-8"))
        for rows in ([[]], [[database_doc()], None]):
            store, _, _ = store_with_db(rows)
            self.assertEqual(store.check(facts)["status"], "insufficient_information")

    def test_limited_context_is_explicit(self):
        store, _, _ = store_with_db([
            [database_doc()], [{"chunk_id": "one"}], [{"chunk_id": "one"}],
            [chunk_row("one", ["two", "three"])], []])
        store.settings = replace(store.settings, policy_context_limit=1)
        self.assertTrue(store.search("住宿")["context_incomplete"])

    def test_real_demo_rules_reference_existing_native_clauses(self):
        # The document is an explicit local fixture, not loaded by runtime PolicyStore.
        path = ROOT / "artifacts/agent有限公司差旅费报销管理制度_模拟修订版_V1.1.docx"
        if not path.exists():
            self.skipTest("Optional user-owned Word fixture absent")
        parsed = parse_docx(path, company_id="test_co", policy_key="travel",
                            version="test", effective_from=date(2026, 9, 7))
        clauses = {c.clause for c in parsed.chunks}
        rules = ReviewedRules.model_validate_json((ROOT / "policies/demo-v1.1.rules.json").read_text("utf-8"))
        required = {c for group in rules.source_clauses.values() for c in group}
        self.assertTrue(required <= clauses)
        self.assertEqual(sum(c.kind == "table_row" for c in parsed.chunks), 9)
        cases = json.loads((ROOT / "policies/demo-retrieval-cases.json").read_text("utf-8"))
        self.assertEqual(len(cases), 30)
        self.assertTrue(all(set(case["clauses"]) <= clauses for case in cases))

    def test_config_check_is_offline_and_redacted(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "nonexistent"
            with patch.dict(os.environ, {"TRAVEL_LOAD_DOTENV": "0"}, clear=True):
                config = Settings.from_env(path)
                self.assertIn("TRAVEL_EMBEDDING_DIMENSIONS", config.missing_policy_config())
                self.assertEqual(config.policy_backend, "local")
        self.assertNotIn("synthetic-secret", repr(settings()))
        self.assertNotIn("test-secret", repr(settings()))


if __name__ == "__main__":
    unittest.main()
