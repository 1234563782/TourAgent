"""Tests for approval form and reimbursement cheat sheet generator."""

import tempfile
import unittest
from pathlib import Path

from docx import Document

from business_travel.approval_forms import (
    calculate_reimbursement_summary,
    export_approval_sheet_docx,
    render_reimbursement_markdown,
)


class ApprovalFormTests(unittest.TestCase):
    def test_reimbursement_calculation_for_ordinary_and_special_regions(self):
        # 1. Ordinary region (Beijing, 3 days)
        facts = {
            "origin": "合肥",
            "destination": "北京",
            "departure_date": "2026-11-01",
            "return_date": "2026-11-03",
        }
        summary = calculate_reimbursement_summary(facts, transport_cost=1200.0, lodging_cost=800.0)
        self.assertEqual(summary["days"], 3)
        self.assertEqual(summary["daily_meal"], 100.0)
        self.assertEqual(summary["total_meal"], 300.0)
        self.assertEqual(summary["daily_local_transport"], 80.0)
        self.assertEqual(summary["total_local_transport"], 240.0)
        self.assertEqual(summary["total_budget"], 1200.0 + 800.0 + 300.0 + 240.0)

        # 2. Special region (Tibet Lhasa, 120 RMB/day meal allowance)
        tibet_facts = {
            "origin": "合肥",
            "destination": "西藏拉萨",
            "departure_date": "2026-11-01",
            "return_date": "2026-11-02",
        }
        tibet_summary = calculate_reimbursement_summary(tibet_facts)
        self.assertEqual(tibet_summary["days"], 2)
        self.assertEqual(tibet_summary["daily_meal"], 120.0)
        self.assertEqual(tibet_summary["total_meal"], 240.0)

    def test_render_reimbursement_markdown(self):
        facts = {
            "user_name": "张教授",
            "origin": "合肥",
            "destination": "上海",
            "departure_date": "2026-11-10",
            "return_date": "2026-11-11",
            "purpose": "参加学术研讨会",
        }
        text = render_reimbursement_markdown(facts, transport_cost=600.0, lodging_cost=400.0)
        self.assertIn("张教授", text)
        self.assertIn("上海", text)
        self.assertIn("伙食补助费（定额包干）", text)
        self.assertIn("市内交通补助费（定额包干）", text)
        self.assertIn("报销总预算预估", text)
        self.assertIn("合肥工业大学差旅报销预审对账单", text)

    def test_export_approval_sheet_docx(self):
        facts = {
            "user_name": "王老师",
            "department": "计算机与信息学院",
            "origin": "合肥",
            "destination": "南京",
            "departure_date": "2026-12-01",
            "return_date": "2026-12-02",
            "purpose": "国家自然科学基金课题交流",
        }
        # In-memory buffer test (bypasses Windows tempdir sandbox permissions)
        buf = export_approval_sheet_docx(facts)
        self.assertTrue(len(buf.getvalue()) > 0)

        # Read back and verify docx content
        doc = Document(buf)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("合肥工业大学出差审批单", full_text)
        self.assertIn("计算机与信息学院", full_text)

        table_text = "\n".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
        self.assertIn("王老师", table_text)
        self.assertIn("国家自然科学基金课题交流", table_text)
        self.assertIn("合肥 - 南京 - 合肥", table_text)

    def test_export_rental_or_driving_approval_sheet_docx(self):
        facts = {
            "user_name": "李老师",
            "department": "机械工程学院",
            "destination": "六安野外实测",
            "purpose": "野外考察自驾",
            "vehicle_detail": "自驾车（车牌皖A-12345）",
        }
        buf = export_approval_sheet_docx(facts, is_rental_or_driving=True, budget_estimate=800.0)
        self.assertTrue(len(buf.getvalue()) > 0)
        doc = Document(buf)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("合肥工业大学租车（自驾车）出差审批表", full_text)
        self.assertIn("注：因自驾车或者租车所引起的安全等问题", full_text)


if __name__ == "__main__":
    unittest.main()
